"""
Fetch historical PVPC electricity prices for 2025 from REData API.
Saves hourly data to an Excel file.

Usage:
    pip install requests openpyxl
    python fetch_prices_2025.py
"""

import requests
import time
import calendar
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, numbers
except ImportError:
    print("Missing dependency. Install with: pip install openpyxl")
    raise SystemExit(1)

API_URL = "https://apidatos.ree.es/en/datos/mercados/precios-mercados-tiempo-real"
PVPC_SERIES_ID = "1001"
SPOT_SERIES_ID = "600"
OUTPUT_FILE = Path(__file__).parent / "pvpc_prices_2025.xlsx"
YEAR = 2025
REQUEST_DELAY = 2  # seconds between requests to be polite


def fetch_month(year: int, month: int) -> list[dict]:
    """Fetch hourly PVPC and spot prices for a given month."""
    last_day = calendar.monthrange(year, month)[1]
    params = {
        "start_date": f"{year}-{month:02d}-01T00:00",
        "end_date": f"{year}-{month:02d}-{last_day}T23:59",
        "time_trunc": "hour",
        "geo_limit": "peninsular",
    }
    headers = {"Accept": "application/json"}

    print(f"  Fetching {year}-{month:02d} ... ", end="", flush=True)
    resp = requests.get(API_URL, params=params, headers=headers, timeout=60)
    resp.raise_for_status()
    data = resp.json()

    # Extract PVPC and spot series from included
    pvpc_values = []
    spot_values = []
    for series in data.get("included", []):
        sid = series.get("id", "")
        if sid == PVPC_SERIES_ID:
            pvpc_values = series.get("attributes", {}).get("values", [])
        elif sid == SPOT_SERIES_ID:
            spot_values = series.get("attributes", {}).get("values", [])

    # Index spot prices by datetime for joining
    spot_by_dt = {}
    for v in spot_values:
        dt_str = v.get("datetime", "")
        spot_by_dt[dt_str] = v.get("value")

    rows = []
    for v in pvpc_values:
        dt_str = v.get("datetime", "")
        try:
            dt = datetime.fromisoformat(dt_str)
        except (ValueError, TypeError):
            continue
        rows.append({
            "datetime": dt,
            "date": dt.strftime("%Y-%m-%d"),
            "hour": dt.hour,
            "pvpc_eur_mwh": v.get("value"),
            "spot_eur_mwh": spot_by_dt.get(dt_str),
        })

    print(f"{len(rows)} hours")
    return rows


def write_excel(all_rows: list[dict], path: Path):
    """Write price data to Excel with formatting."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hourly Prices"

    # Header
    headers = ["Date", "Hour", "PVPC (EUR/MWh)", "PVPC (ct/kWh)", "Spot (EUR/MWh)", "Spot (ct/kWh)"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for i, row in enumerate(all_rows, 2):
        ws.cell(row=i, column=1, value=row["date"])
        ws.cell(row=i, column=2, value=row["hour"])

        pvpc = row["pvpc_eur_mwh"]
        spot = row["spot_eur_mwh"]

        c = ws.cell(row=i, column=3, value=pvpc)
        c.number_format = "0.00"

        # Convert EUR/MWh to ct/kWh (divide by 10)
        c = ws.cell(row=i, column=4, value=round(pvpc / 10, 4) if pvpc is not None else None)
        c.number_format = "0.0000"

        c = ws.cell(row=i, column=5, value=spot)
        c.number_format = "0.00"

        c = ws.cell(row=i, column=6, value=round(spot / 10, 4) if spot is not None else None)
        c.number_format = "0.0000"

    # Auto-width columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    # Daily summary sheet
    ws2 = wb.create_sheet("Daily Summary")
    summary_headers = [
        "Date", "Min PVPC", "Max PVPC", "Avg PVPC",
        "Min Spot", "Max Spot", "Avg Spot", "Avg PVPC (ct/kWh)"
    ]
    for col, h in enumerate(summary_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Group by date
    daily = {}
    for row in all_rows:
        d = row["date"]
        if d not in daily:
            daily[d] = {"pvpc": [], "spot": []}
        if row["pvpc_eur_mwh"] is not None:
            daily[d]["pvpc"].append(row["pvpc_eur_mwh"])
        if row["spot_eur_mwh"] is not None:
            daily[d]["spot"].append(row["spot_eur_mwh"])

    for i, (date, vals) in enumerate(sorted(daily.items()), 2):
        ws2.cell(row=i, column=1, value=date)
        if vals["pvpc"]:
            ws2.cell(row=i, column=2, value=round(min(vals["pvpc"]), 2))
            ws2.cell(row=i, column=3, value=round(max(vals["pvpc"]), 2))
            avg_pvpc = sum(vals["pvpc"]) / len(vals["pvpc"])
            ws2.cell(row=i, column=4, value=round(avg_pvpc, 2))
            ws2.cell(row=i, column=8, value=round(avg_pvpc / 10, 4))
        if vals["spot"]:
            ws2.cell(row=i, column=5, value=round(min(vals["spot"]), 2))
            ws2.cell(row=i, column=6, value=round(max(vals["spot"]), 2))
            ws2.cell(row=i, column=7, value=round(sum(vals["spot"]) / len(vals["spot"]), 2))

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws2.column_dimensions[col[0].column_letter].width = max_len + 2

    # Monthly summary sheet
    ws3 = wb.create_sheet("Monthly Summary")
    monthly_headers = [
        "Month", "Min PVPC", "Max PVPC", "Avg PVPC",
        "Min Spot", "Max Spot", "Avg Spot", "Avg PVPC (ct/kWh)", "Hours"
    ]
    for col, h in enumerate(monthly_headers, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    monthly = {}
    for row in all_rows:
        m = row["date"][:7]  # YYYY-MM
        if m not in monthly:
            monthly[m] = {"pvpc": [], "spot": []}
        if row["pvpc_eur_mwh"] is not None:
            monthly[m]["pvpc"].append(row["pvpc_eur_mwh"])
        if row["spot_eur_mwh"] is not None:
            monthly[m]["spot"].append(row["spot_eur_mwh"])

    for i, (month, vals) in enumerate(sorted(monthly.items()), 2):
        ws3.cell(row=i, column=1, value=month)
        if vals["pvpc"]:
            ws3.cell(row=i, column=2, value=round(min(vals["pvpc"]), 2))
            ws3.cell(row=i, column=3, value=round(max(vals["pvpc"]), 2))
            avg_pvpc = sum(vals["pvpc"]) / len(vals["pvpc"])
            ws3.cell(row=i, column=4, value=round(avg_pvpc, 2))
            ws3.cell(row=i, column=8, value=round(avg_pvpc / 10, 4))
            ws3.cell(row=i, column=9, value=len(vals["pvpc"]))
        if vals["spot"]:
            ws3.cell(row=i, column=5, value=round(min(vals["spot"]), 2))
            ws3.cell(row=i, column=6, value=round(max(vals["spot"]), 2))
            ws3.cell(row=i, column=7, value=round(sum(vals["spot"]) / len(vals["spot"]), 2))

    for col in ws3.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws3.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(path)


def main():
    print(f"Fetching PVPC & spot prices for {YEAR} from REData API")
    print(f"Output: {OUTPUT_FILE}\n")

    all_rows = []
    for month in range(1, 13):
        try:
            rows = fetch_month(YEAR, month)
            all_rows.extend(rows)
        except requests.HTTPError as e:
            print(f"HTTP error: {e}")
        except Exception as e:
            print(f"Error: {e}")

        if month < 12:
            time.sleep(REQUEST_DELAY)

    print(f"\nTotal hours fetched: {len(all_rows)}")

    if all_rows:
        write_excel(all_rows, OUTPUT_FILE)
        print(f"Saved to {OUTPUT_FILE}")

        # Quick stats
        pvpc_prices = [r["pvpc_eur_mwh"] for r in all_rows if r["pvpc_eur_mwh"] is not None]
        if pvpc_prices:
            print(f"\nPVPC Stats (EUR/MWh):")
            print(f"  Min:  {min(pvpc_prices):.2f}")
            print(f"  Max:  {max(pvpc_prices):.2f}")
            print(f"  Avg:  {sum(pvpc_prices)/len(pvpc_prices):.2f}")
    else:
        print("No data fetched.")


if __name__ == "__main__":
    main()
